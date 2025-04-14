"""
Workspace tab UI components for Visual Transcription app
"""
import streamlit as st
import cv2
import numpy as np
from PIL import Image
from streamlit_drawable_canvas import st_canvas
import tempfile
import os
from docx import Document

from image_processing import get_frame_timestamp, seconds_to_timestamp, crop_rectangular, crop_freeform
from transcription import merge_transcripts

def render_workspace_tab():
    """Render the visual transcription workspace tab."""
    # --- Display the video frame and controls only if video is uploaded and available ---
    if st.session_state.get("uploaded", False) and st.session_state.get("video") is not None:
        try:
            # --- Get Current Frame ---
            video_obj = st.session_state.video
            if video_obj is None or not video_obj.isOpened():
                st.error("Video object is not available or not opened. Please load your video in the Media Upload tab.")
                st.session_state.uploaded = False  # Reset the uploaded flag to force re-upload
            else:
                # Add prompt selection radio buttons
                st.subheader("Transcription Settings")
                st.write("Select a prompt category for image description when transcribing:")
                
                # Initialize the prompt category in session state if it doesn't exist
                if "prompt_category" not in st.session_state:
                    st.session_state.prompt_category = "general"
                
                # Create radio buttons for prompt categories
                prompt_category = st.radio(
                    "Prompt Category",
                    ["General Purpose", "STEM", "Humanities & Social Sciences", "Business"],
                    index=0,  # Default to General
                    key="prompt_radio"
                )
                
                # Update the session state based on selection
                category_mapping = {
                    "General Purpose": "general",
                    "STEM": "stem",
                    "Humanities & Social Sciences": "humanities",
                    "Business": "business"
                }
                
                # Update prompt category in session state if changed
                selected_category = category_mapping[prompt_category]
                if st.session_state.prompt_category != selected_category:
                    st.session_state.prompt_category = selected_category
                    
                    # Get the prompt text directly from session state using the prompt's display name
                    if prompt_category in st.session_state:
                        # Use the prompt text that was already loaded by load_all_system_prompts
                        st.session_state['gpt-4o']["prompt"] = st.session_state[prompt_category]
                        st.success(f"Loaded {prompt_category} prompt")
                    else:
                        # If not found in session state, try to load it from file
                        try:
                            with open(f"database/prompts/{selected_category}.json", "r") as json_file:
                                prompt_data = json.load(json_file)
                                st.session_state['gpt-4o']["prompt"] = prompt_data["prompt"].replace(
                                    "%MAX_WORDS%", str(st.session_state["max_words"]))
                                st.success(f"Loaded {prompt_data['name']} prompt")
                        except Exception as e:
                            st.error(f"Error loading prompt: {e}")
                
                # Instead of displaying the full prompt, just show which prompt category is active
                st.markdown(f"### Using {prompt_category} prompt")
                
                st.markdown("---")  # Add separator
                
                video_obj.set(cv2.CAP_PROP_POS_FRAMES, st.session_state.frame_number)
                ret, frame_bgr = video_obj.read()  # Keep original frame in BGR

                if ret:
                    # --- Prepare for Canvas ---
                    # Make a copy for display conversion to avoid modifying original frame_bgr
                    frame_rgb = cv2.cvtColor(frame_bgr.copy(), cv2.COLOR_BGR2RGB)
                    pil_image_bg = Image.fromarray(frame_rgb)

                    # Initialize the crop variable to avoid NameError when checking later
                    current_processed_crop_bgr = None

                    # --- Canvas Mode and Display ---
                    use_rect_mode = st.checkbox("Use Rectangular Crop Mode", value=True, key='crop_mode_checkbox')

                    # Define canvas dimensions (use frame dimensions)
                    canvas_height, canvas_width = frame_bgr.shape[:2]
                    # Optional: Limit max display size for very large videos
                    display_width = min(canvas_width, 1080)
                    display_height = int(display_width * (canvas_height / canvas_width))  # Maintain exact aspect ratio
                    
                    st.write(f"Draw a **{'Rectangle' if use_rect_mode else 'Freeform Shape'}** on the image below.")

                    # Use the canvas_key from session state to force redraw when needed
                    current_canvas_key = f"main_canvas_video_{st.session_state.canvas_key}"
                    try:
                        canvas_result = st_canvas(
                            fill_color="rgba(255, 165, 0, 0.3)",
                            stroke_width=st.session_state.get('stroke_slider', 3), # Use stroke width from settings
                            stroke_color=st.session_state.get('stroke_color', '#00FF00'), # Use stroke color from settings
                            background_color="#eee",
                            background_image=pil_image_bg, # Use PIL image here
                            update_streamlit=True,
                            height=display_height, # Use calculated height to maintain aspect ratio
                            width=display_width,   # Use calculated display size
                            drawing_mode="rect" if use_rect_mode else "freedraw",
                            key=current_canvas_key # Dynamic key that changes to force canvas reset
                        )
                    except Exception as canvas_error:
                        st.error(f"Error rendering canvas: {canvas_error}")
                        canvas_result = None

                    # --- Process Canvas Result (no preview, just processing) ---
                    if canvas_result and canvas_result.json_data is not None and canvas_result.json_data.get("objects"):
                        last_object = canvas_result.json_data["objects"][-1]
                        # Scale factor if canvas size was different from original frame size
                        scale_x = canvas_width / display_width
                        scale_y = canvas_height / display_height

                        if use_rect_mode and last_object["type"] == "rect":
                            # Scale coordinates back to original frame dimensions
                            scaled_rect_data = {
                                'left': last_object['left'] * scale_x,
                                'top': last_object['top'] * scale_y,
                                'width': last_object['width'] * scale_x,
                                'height': last_object['height'] * scale_y,
                            }
                            current_processed_crop_bgr = crop_rectangular(frame_bgr, scaled_rect_data)

                        elif not use_rect_mode and last_object["type"] == "path":
                            # Scale path points back to original frame dimensions
                            original_path_data = []
                            for point_cmd in last_object["path"]:
                                scaled_cmd = [point_cmd[0]] # Keep command
                                # Scale coordinate values
                                for i in range(1, len(point_cmd)):
                                    scaled_cmd.append(point_cmd[i] * (scale_x if i % 2 != 0 else scale_y))
                                original_path_data.append(scaled_cmd)

                            current_processed_crop_bgr = crop_freeform(frame_bgr, original_path_data)
                            
                    # --- Frame Navigation ---
                    st.markdown("---")  # Add separator
                    fps = video_obj.get(cv2.CAP_PROP_FPS)
                    total_frames = st.session_state.total_frames
                    total_duration = total_frames / fps if fps > 0 else 0
                    
                    # Convert current frame to time
                    current_time = st.session_state.frame_number / fps if fps > 0 else 0
                    
                    # Create time-based slider
                    time_input = st.slider(
                        'Select time (seconds)',
                        0.0,
                        max(0.1, total_duration),
                        current_time,
                        step=1.0,  # Increment by 1 second
                        key='time_slider'
                    )
                    
                    # Convert selected time back to frame number
                    new_frame_number = int(time_input * fps) if fps > 0 else 0
                    
                    # Update frame number only if time selection changes
                    if new_frame_number != st.session_state.frame_number:
                        st.session_state.frame_number = new_frame_number
                        # Clear transient crop when navigating away
                        try:
                            current_processed_crop_bgr = None
                        except:
                            pass
                        # Force rerun to update the display with the new frame
                        st.experimental_rerun()

                    # Display both time and frame information
                    st.write(f"Current Time: {current_time:.2f}s (Frame: {st.session_state.frame_number})")
                    
                    # --- Navigation and Save Buttons ---
                    col1, col2, col3 = st.columns([1, 2, 1])
                    with col1:
                        if st.button(f'⬅️ Back {st.session_state.frame_increment} Frame{"s" if st.session_state.frame_increment > 1 else ""}'):
                            if st.session_state.frame_number > 0:
                                # Calculate new frame number with bounds checking
                                new_frame = st.session_state.frame_number - st.session_state.frame_increment
                                # Ensure we don't go below 0
                                st.session_state.frame_number = max(0, new_frame)
                                # Clear transient crop when navigating away
                                try:
                                    current_processed_crop_bgr = None
                                except:
                                    pass
                                st.experimental_rerun()

                    with col2:
                        # Save Button
                        if st.button('💾 Save Frame (Crop if Drawn)'):
                            # Re-get the frame to ensure it's the one displayed
                            video_obj = st.session_state.video
                            if video_obj and video_obj.isOpened():
                                video_obj.set(cv2.CAP_PROP_POS_FRAMES, st.session_state.frame_number)
                                ret_save, frame_bgr_save = video_obj.read()

                                if ret_save:
                                    saved_image_data_rgb = None
                                    is_cropped_flag = False

                                    # Check if there's a crop to use
                                    if current_processed_crop_bgr is not None and current_processed_crop_bgr.size > 0:
                                        # Save the cropped version (convert to RGB)
                                        saved_image_data_rgb = cv2.cvtColor(current_processed_crop_bgr, cv2.COLOR_BGR2RGB)
                                        is_cropped_flag = True
                                        st.success(f"Saving **cropped** frame {st.session_state.frame_number}")
                                    else:
                                        # Save the full frame version (convert to RGB)
                                        saved_image_data_rgb = cv2.cvtColor(frame_bgr_save, cv2.COLOR_BGR2RGB)
                                        is_cropped_flag = False
                                        st.success(f"Saving **full** frame {st.session_state.frame_number}")

                                    # Store in session state with frame info
                                    try:
                                        # Create a copy of the numpy array
                                        frame_copy = saved_image_data_rgb.copy()
                                        
                                        # Store in session state with clear structure
                                        st.session_state.saved_frames[st.session_state.frame_number] = {
                                            'frame': frame_copy, # Store RGB numpy array copy
                                            'frame_number': st.session_state.frame_number,
                                            'is_cropped': is_cropped_flag,
                                            'has_visual_transcripts': False,
                                            'getting_visual_transcripts': False,
                                            'visual_transcripts': None
                                        }
                                        
                                        # Increment the canvas key to force a redraw/clear of the canvas
                                        st.session_state.canvas_key += 1
                                        
                                        # Clear the crop preview after saving
                                        current_processed_crop_bgr = None
                                        
                                        st.success(f"Saved frame {st.session_state.frame_number}")
                                        st.experimental_rerun()
                                    except Exception as save_error:
                                        st.error(f"Error saving frame: {save_error}")
                                else:
                                    st.error('Could not capture the frame to save.')
                            else:
                                st.error("Video is not available. Please load your video in the Media Upload tab.")

                    with col3:
                        if st.button(f'➡️ Forward {st.session_state.frame_increment} Frame{"s" if st.session_state.frame_increment > 1 else ""}'):
                            if st.session_state.frame_number < st.session_state.total_frames - 1:
                                # Calculate new frame number with bounds checking
                                new_frame = st.session_state.frame_number + st.session_state.frame_increment
                                # Ensure we don't go beyond the last frame
                                st.session_state.frame_number = min(st.session_state.total_frames - 1, new_frame)
                                # Clear transient crop when navigating away
                                try:
                                    current_processed_crop_bgr = None
                                except:
                                    pass
                                st.experimental_rerun()
                else:
                    st.error(f'Could not read frame {st.session_state.frame_number}. End of video or error.')
                    
                    # Add a button to allow resetting the video
                    if st.button("Reset Video"):
                        if "video" in st.session_state:
                            try:
                                st.session_state.video.release()
                            except:
                                pass
                        st.session_state.video = None
                        st.session_state.uploaded = False
                        st.experimental_rerun()
        except Exception as e:
            st.error(f"Error accessing or processing video: {e}")
            st.info("Please try uploading your video again in the Media Upload tab.")
            # Reset video-related session state to force clean reload
            st.session_state.uploaded = False
            if "video" in st.session_state:
                try:
                    st.session_state.video.release()
                except:
                    pass
                st.session_state.video = None
    else:
        # Display a message when no video is uploaded (removed placeholder image)
        st.info("Please upload a video in the Media Upload tab to begin.")
        
        # Add a button to allow resetting the session state if needed
        if st.session_state.get("uploaded", False) or st.session_state.get("video") is not None:
            if st.button("Reset Video State"):
                # Clean up video if it exists
                if "video" in st.session_state and st.session_state.video is not None:
                    try:
                        st.session_state.video.release()
                    except:
                        pass
                # Reset all video-related state variables
                st.session_state.video = None
                st.session_state.uploaded = False
                st.session_state.frame_number = 0
                st.experimental_rerun()
    
    # --- Display Final Transcript Information (moved into the workspace tab) ---
    st.markdown("---")
    st.subheader("Combined Transcript (Preview)")
    
    # Use our new merge_transcripts function to get combined chronological transcript
    if st.session_state.audio_transcript:
        st.write(st.session_state.audio_transcript)
    else:
        # Get merged transcripts in chronological order
        merged_transcripts = merge_transcripts()
        
        if merged_transcripts:
            st.write("Showing all transcripts in chronological order:")
            
            # Display merged transcripts
            for entry in merged_transcripts:
                if len(entry) == 3:  # Audio entry (timestamp, text, "audio")
                    timestamp, text, source = entry
                    formatted_time = seconds_to_timestamp(timestamp)
                    st.write(f"**[{formatted_time}]** {text}")
                else:  # Visual entry (timestamp, text, "visual", frame_number)
                    timestamp, text, source, frame_number = entry
                    formatted_time = seconds_to_timestamp(timestamp)
                    st.write(f"**[{formatted_time} - Frame {frame_number}]** 🖼️ *{text}*")
                
            # Add a download option for the combined transcript
            if st.button("Update Combined Transcript Document"):
                # Create a new document with merged transcripts
                doc = Document()
                doc.add_heading("Combined Visual and Audio Transcript", level=1)
                doc.add_paragraph("Timestamps are shown in HH:MM:SS format.")
                for entry in merged_transcripts:
                    if len(entry) == 3:  # Audio entry
                        timestamp, text, source = entry
                        formatted_time = seconds_to_timestamp(timestamp)
                        doc.add_paragraph(f"[{formatted_time}] {text}")
                    else:  # Visual entry
                        timestamp, text, source, frame_number = entry
                        formatted_time = seconds_to_timestamp(timestamp)
                        doc.add_paragraph(f"[{formatted_time} - Frame {frame_number}] {text}")
                
                # Save and offer for download
                temp_doc_path = tempfile.NamedTemporaryFile(delete=False, suffix=".docx").name
                doc.save(temp_doc_path)
                with open(temp_doc_path, "rb") as doc_file:
                    st.download_button(
                        "Download Combined Transcript", 
                        doc_file, 
                        file_name="combined_transcript.docx", 
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
        else:
            st.info("No transcript data available. Please upload an SRT file in the Media Upload tab or add frame transcriptions.") 